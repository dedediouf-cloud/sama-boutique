from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def set_cell_shading(cell, fill):
    """Set background shading for a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)

def add_heading_custom(doc, text, level=1):
    """Add a styled heading."""
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = RGBColor(0xB7, 0x6E, 0x79)  # rose gold
        if level == 1:
            run.font.size = Pt(24)
        elif level == 2:
            run.font.size = Pt(18)
        else:
            run.font.size = Pt(14)
    return heading

def add_tip(doc, text):
    """Add a tip/callout box."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.right_indent = Inches(0.2)
    run = p.add_run("💡 Astuce : " + text)
    run.italic = True
    run.font.color.rgb = RGBColor(0x6B, 0x5B, 0x55)
    p.paragraph_format.space_after = Pt(12)

def add_warning(doc, text):
    """Add a warning box."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.right_indent = Inches(0.2)
    run = p.add_run("⚠️ Important : " + text)
    run.bold = True
    run.font.color.rgb = RGBColor(0xB7, 0x6E, 0x79)
    p.paragraph_format.space_after = Pt(12)

def add_step(doc, number, text):
    """Add a numbered step."""
    p = doc.add_paragraph(style='List Number')
    p.add_run(text).font.color.rgb = RGBColor(0x4A, 0x3F, 0x3A)

def create_documentation():
    doc = Document()
    
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x4A, 0x3F, 0x3A)
    
    # Title page
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run('SamaBoutique')
    run.font.size = Pt(48)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0xB7, 0x6E, 0x79)
    
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run('Guide Utilisateur')
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor(0x6B, 0x5B, 0x55)
    
    version = doc.add_paragraph()
    version.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = version.add_run('Version 1.1 — Juillet 2026')
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x6B, 0x5B, 0x55)
    
    doc.add_page_break()
    
    # Table of contents placeholder
    add_heading_custom(doc, 'Table des matières', level=1)
    toc_items = [
        '1. Introduction',
        '2. Connexion et rôles',
        '3. Tableau de bord',
        '4. Gestion du stock',
        '5. Inventaire',
        '6. Caisse et ventes',
        '7. Clients et fidélité',
        '8. Réservations et commandes',
        '9. Livraisons',
        '10. Promotions',
        '11. Fournisseurs',
        '12. Employés',
        '13. Statistiques',
        '14. Catalogue en ligne et WhatsApp',
        '15. Guide Super Admin',
        '16. Abonnements mensuels / annuels et parrainage',
        '17. FAQ et dépannage'
    ]
    for item in toc_items:
        p = doc.add_paragraph(item, style='List Bullet')
        p.paragraph_format.space_after = Pt(4)
    
    doc.add_page_break()
    
    # 1. Introduction
    add_heading_custom(doc, '1. Introduction', level=1)
    doc.add_paragraph(
        'SamaBoutique est une application de gestion pour petits commerçants, boutiques et vendeurs. '
        'Elle permet de gérer le stock, les ventes, les clients, les livraisons et les commandes en ligne via WhatsApp.'
    )
    doc.add_paragraph(
        'Avec SamaBoutique, vous pouvez :'
    )
    benefits = [
        'Gérer vos produits et votre stock en temps réel',
        'Vendre en boutique avec un panier et une caisse simples',
        'Générer des factures PDF automatiquement',
        'Suivre vos clients et leurs points de fidélité',
        'Recevoir des commandes et réservations via WhatsApp',
        'Gérer les livraisons et les fournisseurs',
        'Créer des promotions et des codes promo',
        'Consulter des statistiques de vente'
    ]
    for benefit in benefits:
        doc.add_paragraph(benefit, style='List Bullet')
    
    add_tip(doc, 'SamaBoutique fonctionne dans un navigateur web (Chrome, Firefox, Edge). Aucune installation sur votre téléphone n\'est nécessaire.')
    
    # 2. Connexion et rôles
    add_heading_custom(doc, '2. Connexion et rôles', level=1)
    add_heading_custom(doc, '2.1 Se connecter', level=2)
    steps = [
        'Ouvrez votre navigateur et allez sur l\'adresse de votre boutique.',
        'Cliquez sur la page de connexion.',
        'Entrez votre email et votre mot de passe.',
        'Cliquez sur « Se connecter ».'
    ]
    for i, step in enumerate(steps, 1):
        add_step(doc, i, step)
    
    add_heading_custom(doc, '2.2 Les rôles', level=2)
    doc.add_paragraph('SamaBoutique propose trois rôles :')
    
    # Roles table
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Light Grid Accent 1'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Rôle'
    hdr_cells[1].text = 'Qui ?'
    hdr_cells[2].text = 'Ce qu\'il peut faire'
    
    roles = [
        ('Super Admin', 'Propriétaire de la plateforme', 'Créer des boutiques, gérer tous les accès'),
        ('Admin', 'Gérant de la boutique', 'Tout gérer dans sa boutique : stock, ventes, employés, promotions...'),
        ('Vendeur', 'Employé de la boutique', 'Passer des ventes, consulter les clients et les réservations')
    ]
    for role, who, what in roles:
        row_cells = table.add_row().cells
        row_cells[0].text = role
        row_cells[1].text = who
        row_cells[2].text = what
    
    add_warning(doc, 'Ne partagez jamais votre mot de passe. Chaque employé doit avoir son propre compte.')
    
    # 3. Tableau de bord
    doc.add_page_break()
    add_heading_custom(doc, '3. Tableau de bord', level=1)
    doc.add_paragraph(
        'Le tableau de bord s\'affiche après la connexion. Il donne un aperçu rapide de votre activité :'
    )
    dashboard_items = [
        'Ventes du jour : nombre de ventes et montant total',
        'Stock faible : produits à réapprovisionner',
        'Réservations en attente : commandes clients à traiter',
        'Top produits : articles les plus vendus'
    ]
    for item in dashboard_items:
        doc.add_paragraph(item, style='List Bullet')
    add_tip(doc, 'Cliquez sur une carte du tableau de bord pour aller vers la section correspondante.')
    
    # 4. Gestion du stock
    add_heading_custom(doc, '4. Gestion du stock', level=1)
    add_heading_custom(doc, '4.1 Ajouter un produit', level=2)
    steps = [
        'Allez dans le menu « Stock ».',
        'Cliquez sur « Ajouter un produit ».',
        'Remplissez le nom, la catégorie, le prix, la quantité et le seuil d\'alerte.',
        'Ajoutez une description si nécessaire.',
        'Cliquez sur « Enregistrer le produit ».'
    ]
    for i, step in enumerate(steps, 1):
        add_step(doc, i, step)
    
    add_heading_custom(doc, '4.2 Réapprovisionner un produit', level=2)
    steps = [
        'Dans la liste des produits, cliquez sur « Réapprovisionner ».',
        'Indiquez la quantité à ajouter.',
        'Ajoutez une note (fournisseur, numéro de lot, etc.).',
        'Cliquez sur « Ajouter au stock ».'
    ]
    for i, step in enumerate(steps, 1):
        add_step(doc, i, step)
    
    add_tip(doc, 'Le seuil d\'alerte vous prévient quand un produit est presque en rupture. Par défaut, il est de 5 unités.')
    
    # 5. Inventaire
    doc.add_page_break()
    add_heading_custom(doc, '5. Inventaire', level=1)
    doc.add_paragraph(
        'Le module Inventaire permet de faire un comptage physique du stock, de corriger les écarts et de connaître la valeur de votre stock.'
    )

    add_heading_custom(doc, '5.1 Faire un comptage', level=2)
    steps = [
        'Allez dans le menu « Inventaire ».',
        'Cliquez sur l\'onglet « Comptage ».',
        'Recherchez le produit ou filtrez par catégorie.',
        'Dans la colonne « Stock physique », entrez la quantité réelle constatée.',
        'Vérifiez l\'écart affiché.',
        'Cliquez sur « Valider » pour enregistrer le comptage et mettre à jour le stock.'
    ]
    for i, step in enumerate(steps, 1):
        add_step(doc, i, step)

    add_heading_custom(doc, '5.2 Voir l\'historique des inventaires', level=2)
    doc.add_paragraph(
        'Cliquez sur l\'onglet « Historique » pour voir tous les comptages passés. '
        'Chaque ligne indique : la date, le produit, le stock système avant comptage, la quantité comptée et l\'écart.'
    )

    add_heading_custom(doc, '5.3 Valorisation du stock', level=2)
    doc.add_paragraph(
        'Cliquez sur l\'onglet « Valorisation » pour voir :'
    )
    valuation_items = [
        'La valeur totale du stock',
        'Le nombre total d\'articles',
        'Le nombre de produits en stock faible',
        'La valeur du stock par catégorie'
    ]
    for item in valuation_items:
        doc.add_paragraph(item, style='List Bullet')
    add_tip(doc, 'Faites un inventaire régulier, par exemple une fois par mois, pour garder votre stock à jour.')

    # 6. Caisse et ventes
    doc.add_page_break()
    add_heading_custom(doc, '6. Caisse et ventes', level=1)
    add_heading_custom(doc, '6.1 Passer une vente', level=2)
    steps = [
        'Allez dans le menu « Ventes ».',
        'Cliquez sur les produits à gauche pour les ajouter au panier.',
        'Modifiez les quantités avec les boutons + et -.',
        'Sélectionnez un client existant ou laissez « Client de passage ».',
        'Choisissez le mode de paiement : Espèces, Orange Money ou Wave.',
        'Cliquez sur « Valider la vente et facturer ».'
    ]
    for i, step in enumerate(steps, 1):
        add_step(doc, i, step)
    
    add_heading_custom(doc, '6.2 Appliquer une promotion', level=2)
    steps = [
        'Dans le panier, choisissez une promotion dans la liste déroulante.',
        'La remise se calcule automatiquement sur le total.',
        'Validez la vente.'
    ]
    for i, step in enumerate(steps, 1):
        add_step(doc, i, step)
    
    add_heading_custom(doc, '6.3 Paiement mobile (Orange Money / Wave)', level=2)
    doc.add_paragraph(
        'Pour les paiements Orange Money ou Wave, indiquez le numéro de téléphone du client. '
        'Un paiement simulé est initié. En production, un compte marchand est nécessaire pour recevoir de vrais paiements.'
    )
    add_warning(doc, 'Pendant la phase de test, les paiements Orange Money et Wave sont simulés. Ne demandez pas encore de vrais paiements à vos clients.')
    
    add_heading_custom(doc, '6.4 Historique des ventes', level=2)
    doc.add_paragraph(
        'Cliquez sur l\'onglet « Historique » dans la page Ventes pour voir toutes les ventes passées. '
        'Vous pouvez vérifier le statut d\'un paiement mobile en attente.'
    )
    
    # 7. Clients et fidélité
    add_heading_custom(doc, '7. Clients et fidélité', level=1)
    add_heading_custom(doc, '7.1 Ajouter un client', level=2)
    steps = [
        'Allez dans le menu « Clients ».',
        'Cliquez sur « Ajouter un client ».',
        'Remplissez le nom, le téléphone, l\'email et l\'adresse.',
        'Ajoutez une date de naissance et des notes si besoin.',
        'Cliquez sur « Enregistrer le client ».'
    ]
    for i, step in enumerate(steps, 1):
        add_step(doc, i, step)
    
    add_heading_custom(doc, '7.2 Voir le détail d\'un client', level=2)
    doc.add_paragraph(
        'Cliquez sur « Voir détails » pour afficher :'
    )
    details = [
        'Le total dépensé par le client',
        'Le nombre d\'achats',
        'Les points de fidélité',
        'L\'historique des achats',
        'Les réservations du client'
    ]
    for detail in details:
        doc.add_paragraph(detail, style='List Bullet')
    
    add_tip(doc, 'Vous pouvez modifier les points de fidélité et les notes d\'un client depuis sa fiche détail.')
    
    # 8. Réservations et commandes
    doc.add_page_break()
    add_heading_custom(doc, '8. Réservations et commandes', level=1)
    doc.add_paragraph(
        'Les clients peuvent réserver des produits depuis votre catalogue en ligne. '
        'Les réservations apparaissent dans le menu « Réservations ».'
    )
    steps = [
        'Allez dans le menu « Réservations ».',
        'Filtrez par statut : En attente, Confirmées, Annulées.',
        'Cliquez sur « Confirmer » ou « Annuler » selon la disponibilité du produit.',
        'Contactez le client par WhatsApp ou téléphone pour finaliser.'
    ]
    for i, step in enumerate(steps, 1):
        add_step(doc, i, step)
    
    add_tip(doc, 'Une réservation confirmée ne retire pas automatiquement le produit du stock. Pensez à passer une vente officielle ensuite.')
    
    # 9. Livraisons
    add_heading_custom(doc, '9. Livraisons', level=1)
    add_heading_custom(doc, '9.1 Créer une livraison', level=2)
    steps = [
        'Allez dans le menu « Livraisons ».',
        'Cliquez sur « Nouvelle livraison ».',
        'Choisissez une vente existante avec un client.',
        'Remplissez l\'adresse, le téléphone et les notes.',
        'Cliquez sur « Créer la livraison ».'
    ]
    for i, step in enumerate(steps, 1):
        add_step(doc, i, step)
    
    add_heading_custom(doc, '9.2 Suivre une livraison', level=2)
    doc.add_paragraph(
        'Pour chaque livraison, vous pouvez changer le statut : En attente → En préparation → Expédiée → Livrée → Annulée.'
    )
    
    # 10. Promotions
    add_heading_custom(doc, '10. Promotions', level=1)
    add_heading_custom(doc, '10.1 Créer une promotion', level=2)
    steps = [
        'Allez dans le menu « Promotions ».',
        'Cliquez sur « Nouvelle promotion ».',
        'Donnez un nom et un code promo (optionnel).',
        'Choisissez le type : Pourcentage ou Montant fixe.',
        'Indiquez la valeur de la remise.',
        'Définissez un montant minimum d\'achat si nécessaire.',
        'Choisissez les dates de début et de fin.',
        'Cliquez sur « Créer la promotion ».'
    ]
    for i, step in enumerate(steps, 1):
        add_step(doc, i, step)
    
    add_tip(doc, 'Vous pouvez activer ou désactiver une promotion à tout moment avec le bouton de bascule.')
    
    # 11. Fournisseurs
    doc.add_page_break()
    add_heading_custom(doc, '11. Fournisseurs', level=1)
    add_heading_custom(doc, '11.1 Ajouter un fournisseur', level=2)
    steps = [
        'Allez dans le menu « Fournisseurs ».',
        'Cliquez sur « Fournisseur ».',
        'Remplissez le nom, le téléphone, l\'email et l\'adresse.',
        'Cliquez sur « Enregistrer le fournisseur ».'
    ]
    for i, step in enumerate(steps, 1):
        add_step(doc, i, step)
    
    add_heading_custom(doc, '11.2 Passer une commande fournisseur', level=2)
    steps = [
        'Cliquez sur « Commande ».',
        'Choisissez le fournisseur.',
        'Ajoutez les articles à commander.',
        'Indiquez les quantités et les prix unitaires.',
        'Cliquez sur « Créer la commande ».'
    ]
    for i, step in enumerate(steps, 1):
        add_step(doc, i, step)
    
    add_heading_custom(doc, '11.3 Réceptionner une commande', level=2)
    doc.add_paragraph(
        'Quand vous recevez la marchandise, cliquez sur « Reçue ». '
        'Le stock de chaque produit est automatiquement mis à jour.'
    )
    
    # 12. Employés
    add_heading_custom(doc, '12. Employés', level=1)
    doc.add_paragraph('Cette section est réservée aux admins.')
    steps = [
        'Allez dans le menu « Employés ».',
        'Cliquez sur « Ajouter un employé ».',
        'Remplissez le nom, l\'email et le mot de passe.',
        'Choisissez le rôle : Vendeur ou Admin.',
        'Cliquez sur « Enregistrer l\'employé ».'
    ]
    for i, step in enumerate(steps, 1):
        add_step(doc, i, step)
    add_warning(doc, 'Un vendeur ne peut pas créer de nouveaux employés ni modifier les promotions. Seul un admin a tous les droits.')
    
    # 13. Statistiques
    add_heading_custom(doc, '13. Statistiques', level=1)
    doc.add_paragraph(
        'Le menu « Statistiques » affiche des graphiques de votre activité :'
    )
    stats_items = [
        'Chiffre d\'affaires total',
        'Nombre de ventes',
        'Panier moyen',
        'Évolution des ventes sur 7 jours, 30 jours ou 12 mois',
        'Top produits par chiffre d\'affaires',
        'Répartition des statuts de paiement'
    ]
    for item in stats_items:
        doc.add_paragraph(item, style='List Bullet')
    add_tip(doc, 'Changez la période avec les boutons en haut de la page pour voir les tendances à court ou long terme.')
    
    # 14. Catalogue en ligne et WhatsApp
    doc.add_page_break()
    add_heading_custom(doc, '14. Catalogue en ligne et WhatsApp', level=1)
    doc.add_paragraph(
        'Chaque boutique a un catalogue public accessible par un lien unique.'
    )
    add_heading_custom(doc, '14.1 Trouver le lien de votre catalogue', level=2)
    steps = [
        'Connectez-vous à votre compte.',
        'Dans la barre latérale, cliquez sur « Mon catalogue ».',
        'Le lien s\'ouvre dans un nouvel onglet.'
    ]
    for i, step in enumerate(steps, 1):
        add_step(doc, i, step)
    
    add_heading_custom(doc, '14.2 Partager le catalogue', level=2)
    doc.add_paragraph(
        'Partagez le lien de votre catalogue sur WhatsApp, Instagram, Facebook ou par SMS. '
        'Vos clients verront vos produits avec leurs prix et pourront :'
    )
    catalog_actions = [
        'Commander via WhatsApp',
        'Acheter maintenant avec Orange Money ou Wave (simulé en test)',
        'Réserver un produit'
    ]
    for action in catalog_actions:
        doc.add_paragraph(action, style='List Bullet')
    
    add_tip(doc, 'Votre numéro WhatsApp doit être renseigné dans votre profil boutique pour que les commandes WhatsApp fonctionnent.')
    
    # 15. Guide Super Admin
    add_heading_custom(doc, '15. Guide Super Admin', level=1)
    doc.add_paragraph(
        'Le Super Admin gère plusieurs boutiques sur la même plateforme SamaBoutique.'
    )
    add_heading_custom(doc, '15.1 Se connecter en Super Admin', level=2)
    doc.add_paragraph(
        'Allez sur `/superadmin/login` et connectez-vous avec vos identifiants Super Admin.'
    )
    add_heading_custom(doc, '15.2 Tableau de bord de suivi', level=2)
    doc.add_paragraph(
        'Le tableau de bord Super Admin donne une vue d\'ensemble de toutes les boutiques :'
    )
    dashboard_items = [
        'Boutiques totales, actives et bloquées',
        'Nombre total de produits, clients et ventes',
        'Recurrence mensuelle attendue (MRR) des abonnements',
        'Nombre d\'abonnements payés et d\'abonnements en retard',
        'Liste détaillée avec statut, échéance, montant et actions rapides'
    ]
    for item in dashboard_items:
        doc.add_paragraph(item, style='List Bullet')
    add_tip(doc, 'Utilisez le tableau de bord pour repérer rapidement les boutiques en retard de paiement.')

    add_heading_custom(doc, '15.3 Abonnements mensuels et annuels', level=2)
    doc.add_paragraph(
        'Chaque boutique peut être configurée en abonnement mensuel ou annuel. '
        'L\'abonnement annuel applique une réduction configurable dans les paramètres globaux.'
    )
    steps = [
        'Dans le tableau de bord, consultez la colonne « Intervalle » pour voir si la boutique est mensuelle ou annuelle.',
        'La colonne « Abonnement » affiche le tarif mensuel de base et, pour les abonnements annuels, le montant annuel après réduction.',
        'Quand la boutique vous paie, cliquez sur « Payer ». L\'échéance est reculée d\'un mois ou d\'un an selon l\'intervalle choisi.',
        'Pour changer l\'intervalle d\'une boutique, cliquez sur « Changer » dans la colonne « Intervalle ».',
        'Si la boutique ne paie pas à temps, elle est marquée « En retard ». Vous pouvez la bloquer si nécessaire.'
    ]
    for i, step in enumerate(steps, 1):
        add_step(doc, i, step)
    add_warning(doc, 'Le bouton « Payer » est manuel : il enregistre que vous avez reçu le paiement. Il ne déclenche pas de vrai prélèvement bancaire.')

    add_heading_custom(doc, '15.4 Paramètres globaux d\'abonnement', level=2)
    doc.add_paragraph(
        'Le panneau « Paramètres globaux » du Super Admin permet de configurer les règles commerciales :'
    )
    global_items = [
        'Réduction appliquée aux abonnements annuels (en %)',
        'Nombre de mois gratuits offerts au parrain pour chaque filleul abonné en annuel',
        'Montant mensuel par défaut proposé lors de la création d\'une boutique'
    ]
    for item in global_items:
        doc.add_paragraph(item, style='List Bullet')
    add_tip(doc, 'Ces paramètres sont utilisés automatiquement lors du paiement d\'un abonnement annuel et du calcul des récompenses de parrainage.')

    add_heading_custom(doc, '15.5 Créer une boutique', level=2)
    steps = [
        'Dans le tableau de bord Super Admin, cliquez sur « Créer une boutique ».',
        'Remplissez le nom du gérant, l\'email et le mot de passe.',
        'Indiquez le nom de la boutique et le slug (identifiant dans l\'URL).',
        'Renseignez le numéro WhatsApp et choisissez l\'intervalle de facturation (mensuel ou annuel).',
        'Cliquez sur « Créer la boutique ».'
    ]
    for i, step in enumerate(steps, 1):
        add_step(doc, i, step)
    
    add_heading_custom(doc, '15.6 Bloquer ou débloquer une boutique', level=2)
    doc.add_paragraph(
        'Le Super Admin peut bloquer une boutique en cas de non-paiement d\'abonnement ou de problème. '
        'Une boutique bloquée ne peut plus se connecter. Ses employés non plus ne peuvent pas se connecter.'
    )
    steps = [
        'Dans la liste des boutiques, cliquez sur « Bloquer ».',
        'Confirmez l\'action.',
        'Le statut de la boutique passe à « Bloquée ».',
        'Pour réactiver la boutique, cliquez sur « Débloquer ».'
    ]
    for i, step in enumerate(steps, 1):
        add_step(doc, i, step)
    add_warning(doc, 'Une boutique bloquée ne pourra plus ouvrir de session. Pensez à prévenir le gérant avant de bloquer.')

    add_heading_custom(doc, '15.7 Promotions sur les abonnements', level=2)
    doc.add_paragraph(
        'Le Super Admin peut lancer des campagnes promotionnelles pour les abonnements : remise en pourcentage, remise en FCFA ou mois gratuits.'
    )
    steps = [
        'Dans la section « Campagnes promotionnelles abonnements », cliquez sur « Nouvelle promotion ».',
        'Donnez un nom et un code promo (optionnel).',
        'Choisissez le type d\'avantage : remise en %, remise en FCFA ou mois gratuits.',
        'Cliquez sur « Créer la promotion ».',
        'Quand vous marquez un abonnement comme payé, saisissez le code promo pour appliquer la remise.'
    ]
    for i, step in enumerate(steps, 1):
        add_step(doc, i, step)
    add_tip(doc, 'Le tableau de bord affiche le chiffre d\'affaires réel des abonnements payés (après remise).')

    add_heading_custom(doc, '15.8 Langue de l\'application', level=2)
    doc.add_paragraph(
        'SamaBoutique propose une interface en français et en anglais.'
    )
    steps = [
        'Cliquez sur le bouton FR / EN dans la barre latérale ou dans l\'en-tête Super Admin.',
        'L\'interface bascule immédiatement dans la langue choisie.',
        'Le choix est mémorisé pour votre prochaine connexion.'
    ]
    for i, step in enumerate(steps, 1):
        add_step(doc, i, step)
    add_tip(doc, 'Vous pouvez changer de langue à tout moment sans vous déconnecter.')

    add_heading_custom(doc, '15.9 Supprimer une boutique', level=2)
    doc.add_paragraph(
        'La suppression supprime définitivement la boutique et toutes ses données : produits, clients, ventes, employés, fournisseurs, commandes, promotions et inventaires.'
    )
    steps = [
        'Dans la liste des boutiques, cliquez sur « Supprimer ».',
        'Lisez l\'avertissement et confirmez.',
        'La boutique est supprimée définitivement.'
    ]
    for i, step in enumerate(steps, 1):
        add_step(doc, i, step)
    add_warning(doc, 'La suppression est irréversible. Ne supprimez une boutique que si vous en avez l\'autorisation expresse du gérant.')

    add_heading_custom(doc, '15.10 Gérer les boutiques', level=2)
    doc.add_paragraph(
        'Le Super Admin peut :'
    )
    admin_actions = [
        'Voir toutes les boutiques créées',
        'Réinitialiser le mot de passe d\'une boutique',
        'Bloquer ou débloquer une boutique',
        'Supprimer une boutique',
        'Lancer des promotions sur les abonnements',
        'Voir le catalogue public de chaque boutique',
        'Consulter le nombre de produits, clients, ventes et le CA des abonnements par boutique',
        'Changer l\'intervalle de facturation d\'une boutique',
        'Consulter le code de parrainage et copier le lien de parrainage d\'une boutique'
    ]
    for action in admin_actions:
        doc.add_paragraph(action, style='List Bullet')
    
    add_warning(doc, 'Le compte de démonstration Super Admin (superadmin@boutique.com / demo123) doit être changé en production pour des raisons de sécurité.')
    
    # 16. Abonnements mensuels / annuels et parrainage
    doc.add_page_break()
    add_heading_custom(doc, '16. Abonnements mensuels / annuels et parrainage', level=1)
    doc.add_paragraph(
        'Cette section résume les nouvelles règles d\'abonnement et le système de parrainage de SamaBoutique.'
    )

    add_heading_custom(doc, '16.1 Choisir un abonnement annuel ou mensuel', level=2)
    doc.add_paragraph(
        'Lors de l\'inscription ou de la création par le Super Admin, chaque boutique choisit un intervalle de facturation :'
    )
    interval_items = [
        'Mensuel : tarif de base, échéance fixée au 5 du mois suivant.',
        'Annuel : tarif de base × 12 moins la réduction globale (ex: 10% de remise). L\'échéance est reculée d\'un an.'
    ]
    for item in interval_items:
        doc.add_paragraph(item, style='List Bullet')
    add_tip(doc, 'Le Super Admin peut changer l\'intervalle d\'une boutique à tout moment.')

    add_heading_custom(doc, '16.2 Système de parrainage', level=2)
    doc.add_paragraph(
        'Chaque boutique dispose d\'un code de parrainage unique. Le lien de parrainage est de la forme :'
    )
    doc.add_paragraph('https://votre-domaine.com/register?ref=CODE_DU_PARRAIN', style='List Bullet')
    steps = [
        'Dans le tableau de bord Super Admin, consultez la colonne « Parrainage » pour voir le code, le nombre de filleuls et les récompenses reçues.',
        'Cliquez sur « Copier le lien » pour partager le lien d\'inscription d\'une boutique.',
        'Quand un nouveau client s\'inscrit avec ce lien et choisit un abonnement annuel, son parrain reçoit automatiquement 1 mois gratuit.',
        'Le nombre de mois gratuits est configurable dans les paramètres globaux.'
    ]
    for i, step in enumerate(steps, 1):
        add_step(doc, i, step)
    add_warning(doc, 'La récompense de parrainage n\'est accordée que si le nouveau client s\'abonne en annuel.')

    add_heading_custom(doc, '16.3 Depuis la page d\'inscription', level=2)
    doc.add_paragraph(
        'Un nouveau client peut s\'inscrire avec un code de parrainage pré-rempli depuis l\'URL. '
        'Il choisit son intervalle de facturation et, s\'il opte pour l\'annuel, le parrain reçoit la récompense.'
    )
    add_tip(doc, 'Partagez votre lien de parrainage sur WhatsApp, Instagram ou par SMS pour recruter de nouvelles boutiques.')
    
    # 17. FAQ et dépannage
    doc.add_page_break()
    add_heading_custom(doc, '17. FAQ et dépannage', level=1)
    
    faqs = [
        ('Je ne reçois pas les commandes WhatsApp', 'Vérifiez que votre numéro WhatsApp est bien renseigné dans votre profil boutique. Le catalogue utilise ce numéro pour envoyer les messages.'),
        ('Mon compte est bloqué', 'Votre abonnement mensuel ou annuel est peut-être en retard ou le Super Admin a suspendu votre boutique. Contactez-le pour régulariser la situation et débloquer votre compte.'),
        ('Le paiement Orange Money / Wave ne marche pas', 'En phase de test, les paiements sont simulés. Pour un vrai déploiement, il faut un compte marchand chez Orange Money ou Wave.'),
        ('Un produit n\'apparaît pas dans le catalogue', 'Vérifiez que le produit est bien créé et que sa quantité en stock est supérieure à 0.'),
        ('Je ne vois pas le menu Employés / Promotions / Fournisseurs', 'Ces menus sont réservés aux admins. Si vous êtes vendeur, demandez à votre gérant de vous donner le rôle admin.'),
        ('J\'ai oublié mon mot de passe', 'Contactez le Super Admin ou l\'admin de votre boutique pour qu\'il réinitialise votre mot de passe.'),
        ('L\'application est lente', 'Vérifiez votre connexion internet. Si le problème persiste, rechargez la page ou contactez le support.')
    ]
    
    for question, answer in faqs:
        add_heading_custom(doc, question, level=3)
        doc.add_paragraph(answer)
    
    add_tip(doc, 'Pour toute question supplémentaire, contactez le support de SamaBoutique.')
    
    # Footer
    doc.add_paragraph()
    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run('SamaBoutique — Votre boutique, simplifiée.')
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x6B, 0x5B, 0x55)
    run.italic = True
    
    # Save
    doc.save('SamaBoutique-Guide-Utilisateur.docx')
    print('Documentation générée : SamaBoutique-Guide-Utilisateur.docx')

if __name__ == '__main__':
    create_documentation()
